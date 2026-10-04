import io
from dataclasses import dataclass

from docx import Document as DocxDocument
from pypdf import PdfReader


@dataclass
class ExtractedPage:
    page_number: int
    text: str
    start_offset: int
    end_offset: int


@dataclass
class ExtractedDocument:
    text: str
    pages: list[ExtractedPage]


def extract_document(file_content: bytes, file_type: str) -> ExtractedDocument:
    if file_type == "pdf":
        return extract_pdf(file_content)

    if file_type == "docx":
        return extract_docx(file_content)

    raise ValueError(f"Unsupported file type: {file_type}")


def extract_pdf(file_content: bytes) -> ExtractedDocument:
    reader = PdfReader(io.BytesIO(file_content))

    pages = []
    full_text = []
    current_offset = 0

    for page_number, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""

        if full_text:
            current_offset += 1

        start_offset = current_offset
        end_offset = start_offset + len(page_text)

        pages.append(
            ExtractedPage(
                page_number=page_number,
                text=page_text,
                start_offset=start_offset,
                end_offset=end_offset,
            )
        )

        full_text.append(page_text)
        current_offset = end_offset

    text = "\n".join(full_text)

    return ExtractedDocument(
        text=text,
        pages=pages,
    )

def extract_docx(file_content: bytes) -> ExtractedDocument:
    document = DocxDocument(io.BytesIO(file_content))

    paragraphs = [
        paragraph.text
        for paragraph in document.paragraphs
    ]

    text = "\n".join(paragraphs)

    page = ExtractedPage(
        page_number=1,
        text=text,
        start_offset=0,
        end_offset=len(text),
    )

    return ExtractedDocument(
        text=text,
        pages=[page],
    )