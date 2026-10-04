import io
import uuid
from pathlib import Path

from docx import Document as DocxDocument
from fastapi import UploadFile
from pypdf import PdfReader
from sqlalchemy.orm import Session

from app.enums.document_status import DocumentStatus
from app.utils.app_exceptions import AppException
from app.models.document_model import Document


UPLOAD_DIR = Path("uploads")

ALLOWED_EXTENSIONS = {".pdf", ".docx"}

MIN_TEXT_LENGTH = 100


def save_document(file: UploadFile, db: Session) -> Document:

    extension = Path(file.filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise AppException(
            "Only PDF and DOCX files are supported.",
            400
        )

    try:
        file_content = file.file.read()
    except Exception as e:
        print(f"Error while reading uploaded file: {e}")
        raise AppException(
            "Failed to read uploaded file.",
            400
        )

    try:
        verify_document(file_content, extension)
    except AppException:
        raise
    except Exception as e:
        print(f"Error while verifying document: {e}")
        raise AppException(
            "The uploaded file is invalid or corrupted.",
            400
        )

    document_id = uuid.uuid4()

    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

    storage_filename = f"{document_id}{extension}"
    storage_path = UPLOAD_DIR / storage_filename

    # 5. Save the original file
    try:
        with open(storage_path, "wb") as f:
            f.write(file_content)
    except Exception as e:
        print(f"Error while saving file: {e}")
        raise AppException(
            "Failed to save file.",
            500
        )

    # 6. Create database record
    document = Document(
        id=document_id,
        filename=file.filename,
        file_type=extension.lstrip("."),
        storage_path=str(storage_path),
        status=DocumentStatus.UPLOADED
    )

    try:
        db.add(document)
        db.commit()
        db.refresh(document)
    except Exception as e:
        db.rollback()

        # File was already saved, so remove it if DB insert fails.
        storage_path.unlink(missing_ok=True)

        print(f"Database error while saving document: {e}")

        raise AppException(
            "Failed to save document.",
            500
        )

    return document

def verify_document(file_content: bytes, extension: str) -> None:
    if extension == ".pdf":
        verify_pdf(file_content)

    elif extension == ".docx":
        verify_docx(file_content)

#PDF Verification
def verify_pdf(file_content: bytes) -> None:
    reader = PdfReader(io.BytesIO(file_content))

    full_text = []

    for page in reader.pages:
        page_text = page.extract_text() or ""
        full_text.append(page_text)

    text = "\n".join(full_text)

    # Remove whitespace before checking meaningful content
    pdf_content = "".join(text.split())

    if len(pdf_content) < MIN_TEXT_LENGTH:
        raise AppException(
            "The PDF does not contain enough readable text.",
            400
        )


# docx Verification
def verify_docx(file_content: bytes) -> None:
    document = DocxDocument(io.BytesIO(file_content))

    full_text = []

    for paragraph in document.paragraphs:
        full_text.append(paragraph.text)

    text = "\n".join(full_text)

    docx_content = "".join(text.split())

    if len(docx_content) < MIN_TEXT_LENGTH:
        raise AppException(
            "The DOCX document does not contain enough readable text.",
            400
        )